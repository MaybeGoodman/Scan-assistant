"""Export reviewed blocks to real DOCX after explicit Word confirmation."""
import argparse
import json
import re
from pathlib import Path
from docx import Document
from docx.shared import Cm, Pt
from docx.oxml.ns import qn
from PIL import Image
from export_page import figure_assets, table_html
from content_blocks import parts, check_word_text, reviewed_blocks
from formulas import FormulaError, to_omml, validate_latex, node


def paragraph_layout(blocks, index):
    """Keep adjacent question/option/figure blocks together without linking all pages."""
    block = blocks[index]
    explicit = block.get('layout', {})
    allowed = {'keep_with_next', 'keep_together', 'page_break_before'}
    if not isinstance(explicit, dict) or set(explicit) - allowed:
        raise ValueError('layout supports only keep_with_next, keep_together and page_break_before')
    if any(type(value) is not bool for value in explicit.values()):
        raise ValueError('layout values must be booleans')

    def text(item):
        if item['type'] != 'text':
            return ''
        return ''.join(p['text'] if p['type'] == 'text' else ' ' for p in parts(item)).lstrip()

    def option(item):
        match = re.match(r'^([A-H])[.．、]\s*', text(item))
        return match[1] if match else None

    following = blocks[index + 1] if index + 1 < len(blocks) else None
    keep = False
    if following:
        label, next_label = option(block), option(following)
        question = bool(re.match(r'^\d+[.．、]\s*', text(block)))
        if question and (following['type'] == 'image' or next_label == 'A'):
            keep = True
        elif label and next_label and ord(next_label) == ord(label) + 1:
            keep = True
        elif block['type'] == 'image' and next_label:
            keep = True
        elif label and following['type'] == 'image' and index + 2 < len(blocks):
            after_image = option(blocks[index + 2])
            keep = bool(after_image and ord(after_image) == ord(label) + 1)
    return {'keep_with_next': keep, 'keep_together': True, 'page_break_before': False, **explicit}


def apply_layout(paragraph, settings):
    for key, value in settings.items():
        setattr(paragraph.paragraph_format, key, value)
    paragraph.paragraph_format.widow_control = True


def add_expression(paragraph, expression, records, location, *, display=False, source_location=None):
    kind = expression['type']
    latex = expression.get('latex')
    if not isinstance(latex, str):
        raise ValueError('latex must be a string')
    status = expression.get('status', 'recognized')
    if status not in {'recognized', 'unresolved'}:
        raise ValueError('Expression status must be recognized or unresolved')
    repairs = expression.get('repairs', [])
    if not isinstance(repairs, list):
        raise ValueError('repairs must be a list of source restoration records')
    record = {'location': location, 'source': expression.get('source', source_location),
              'type': kind, 'latex': latex, 'display': display, 'repairs': repairs}
    if 'id' in expression:
        record['id'] = expression['id']
    if 'confidence' in expression:
        record['confidence'] = expression['confidence']
    try:
        if status == 'unresolved':
            raise FormulaError(expression.get('reason', 'Expression could not be reliably recognized'))
        if kind == 'math':
            formula = to_omml(latex, display=display)
            if display:
                wrapper = node('oMathPara', node('oMathParaPr', node('jc', val='center')), formula)
                paragraph._p.append(wrapper)
            else:
                paragraph._p.append(formula)
        elif kind == 'chemistry':
            validate_latex(latex, chemistry=True)
            paragraph.add_run(latex)  # Literal text; never feed chemistry to the math converter.
        else:
            raise ValueError('Unknown expression type: ' + kind)
        record['status'] = 'needs_review' if 'XXX' in latex or repairs else 'converted' if kind == 'math' else 'preserved'
    except FormulaError as exc:
        paragraph.add_run('XXX')
        record.update(status='needs_review', reason=str(exc))
    records.append(record)


def add_content(paragraph, item, records, location, source_location=None):
    for index, part in enumerate(parts(item)):
        if part['type'] == 'text':
            check_word_text(part['text'])
            paragraph.add_run(part['text'])
        else:
            add_expression(paragraph, part, records, f'{location}.runs[{index}]',
                           source_location=item.get('source', source_location))


def export_docx(source, output, *, word_confirmed=False, report_path=None):
    if word_confirmed is not True:
        raise ValueError('Confirm Word output with the user first')
    source, output = Path(source).resolve(), Path(output).resolve()
    if output.suffix.lower() != '.docx':
        raise ValueError('Output must end in .docx')
    if output.exists():
        raise FileExistsError(output)
    report_path = Path(report_path).resolve() if report_path else output.with_suffix('.review.json')
    if report_path in {source, output} or report_path.suffix.lower() != '.json':
        raise ValueError('Review report must be a separate JSON file')
    if report_path.exists():
        raise FileExistsError(report_path)
    blocks = reviewed_blocks(json.loads(source.read_text(encoding='utf-8')))
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin = section.bottom_margin = Cm(2)
    section.left_margin = section.right_margin = Cm(2.2)
    style = document.styles['Normal']
    style.font.name = 'Times New Roman'
    style.font.size = Pt(11)
    style.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), '宋体')
    style.paragraph_format.space_after = Pt(6)
    style.paragraph_format.line_spacing = 1.2
    max_width = int(section.page_width - section.left_margin - section.right_margin)
    max_height = int(section.page_height - section.top_margin - section.bottom_margin - Cm(1))
    records = []
    for block_index, block in enumerate(blocks):
        location = f'blocks[{block_index}]'
        kind = block['type']
        settings = paragraph_layout(blocks, block_index)
        paragraph = None
        if kind == 'text':
            paragraph = document.add_paragraph()
            add_content(paragraph, block, records, location)
        elif kind in {'math', 'chemistry'}:
            paragraph = document.add_paragraph()
            add_expression(paragraph, block, records, location, display=True)
        elif kind == 'table':
            table_html(block)  # Shared grid validation, including blank and merged cells.
            table = document.add_table(rows=block['rows'], cols=block['cols'])
            table.style = 'Table Grid'
            table.autofit = False
            for column in table.columns:
                column.width = int(max_width / block['cols'])
            for cell_index, item in enumerate(block['cells']):
                r, c = item['row'], item['col']
                rs, cs = item.get('rowspan', 1), item.get('colspan', 1)
                cell = table.cell(r, c)
                if rs > 1 or cs > 1:
                    cell = cell.merge(table.cell(r + rs - 1, c + cs - 1))
                cell.text = ''
                add_content(cell.paragraphs[0], item, records, f'{location}.cells[{cell_index}]',
                            block.get('source'))
        elif kind == 'image':
            image, _ = figure_assets(source.parent, block)
            with Image.open(image) as pixels:
                pixels.verify()
            with Image.open(image) as pixels:
                width, height = pixels.size
                dpi = pixels.info.get('dpi', (96, 96))[0]
            # Honour the PNG's own resolution (300 dpi redraws would otherwise be
            # sized as 96 dpi and overflow); missing or implausible values fall back to 96.
            dpi = float(dpi) if isinstance(dpi, (int, float)) and 72 <= dpi <= 1200 else 96.0
            display_width = min(max_width, int(max_height * width / height), int(Cm(width / dpi * 2.54)))
            shape = document.add_picture(str(image), width=display_width)
            paragraph = document.paragraphs[-1]
            if 'redraw' in block:
                from docx.enum.text import WD_ALIGN_PARAGRAPH
                paragraph.alignment=WD_ALIGN_PARAGRAPH.CENTER
            alt = block.get('alt', '')
            if not isinstance(alt, str):
                raise ValueError('Image alt must be a string')
            shape._inline.docPr.set('descr', alt)
        else:
            raise ValueError(f'Unknown block type: {kind}')
        if paragraph is not None:
            apply_layout(paragraph, settings)
    document.core_properties.author = ''
    document.core_properties.last_modified_by = ''
    document.core_properties.comments = ''
    output.parent.mkdir(parents=True, exist_ok=True)
    report_path.parent.mkdir(parents=True, exist_ok=True)
    report = {'schema_version': 1, 'source_file': str(source), 'expressions': records,
              'needs_review': sum(record['status'] == 'needs_review' for record in records),
              'figures':[{'figure':block.get('figure'),'validation':block['redraw']} for block in blocks if 'redraw' in block]}
    # Exclusive creation prevents accidental replacement if another process wrote meanwhile.
    created = []
    try:
        with report_path.open('x', encoding='utf-8') as stream:
            created.append(report_path)
            json.dump(report, stream, ensure_ascii=False, indent=2)
        with output.open('xb') as stream:
            created.append(output)
            document.save(stream)
    except Exception:
        for path in created:
            path.unlink(missing_ok=True)
        raise
    return output


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('source')
    parser.add_argument('--output', required=True)
    parser.add_argument('--word-confirmed', action='store_true',
                        help='Use only after the user explicitly chose Word')
    parser.add_argument('--report', help='Internal review JSON (default: output.review.json)')
    args = parser.parse_args()
    print(export_docx(args.source, args.output, word_confirmed=args.word_confirmed, report_path=args.report))
