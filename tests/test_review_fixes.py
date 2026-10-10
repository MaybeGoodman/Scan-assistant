"""Regressions for reply parsing, Markdown escaping, Word text checks and figure sizing."""
import json
from pathlib import Path
import sys
import tempfile
import unittest
from docx import Document
from docx.shared import Cm
from PIL import Image

sys.path.insert(0, str(Path(__file__).resolve().parent))
from test_figures import document, reviews  # noqa: E402  (also puts scripts on sys.path)
from content_blocks import check_word_text, plain_source  # noqa: E402
from export_docx import export_docx  # noqa: E402
from export_page import export  # noqa: E402
from figure_drawing import check_black_white, render  # noqa: E402
from figure_workflow import complete, generate, reply_mode  # noqa: E402


class ReplyModeTests(unittest.TestCase):
    def test_requirement_examples_still_map(self):
        for reply in ('是', '可以', '需要', '继续', '要', '重绘', '帮我画出来', '图片也处理', '图也生成'):
            self.assertEqual(reply_mode(reply), 'redraw', reply)
        for reply in ('否', '不需要', '不用', '跳过', '只要文字', '不要图片', '图片忽略'):
            self.assertEqual(reply_mode(reply), 'ignore', reply)

    def test_common_short_variants(self):
        for reply in ('好的', '好', '是的', '可以的', '行', '嗯，需要', '好的，需要。', 'OK', 'yes'):
            self.assertEqual(reply_mode(reply), 'redraw', reply)
        for reply in ('不', '不要', '不要了', '不用了吧', '不行', '不是', '好的，不用', '嗯，不要'):
            self.assertEqual(reply_mode(reply), 'ignore', reply)
        for reply in ('不用重绘，原图保留', '不重绘，保留原图', '保留原图'):
            self.assertEqual(reply_mode(reply), 'preserve', reply)

    def test_unclear_or_conflicting_replies_stay_ambiguous(self):
        for reply in ('也许', '要，不要', '随便', '', '  ', '看情况'):
            with self.assertRaises(ValueError, msg=reply):
                reply_mode(reply)


class MarkdownEscapingTests(unittest.TestCase):
    def md(self, **item):
        return plain_source(item, markdown=True)

    def test_line_start_markers_are_literal(self):
        self.assertEqual(self.md(text='1. 下列说法正确的是'), '1\\. 下列说法正确的是')
        self.assertEqual(self.md(text='# 注意'), '\\# 注意')
        self.assertEqual(self.md(text='- 选项\n2) 条件\n+3'), '\\- 选项\\\n2\\) 条件\\\n\\+3')

    def test_inline_markup_html_and_dollars_are_literal(self):
        self.assertEqual(self.md(text='a*b_c <b>x</b> [链接](u) 单价$5'),
                         'a\\*b\\_c \\<b\\>x\\</b\\> \\[链接\\](u) 单价\\$5')

    def test_line_breaks_and_paragraphs(self):
        self.assertEqual(self.md(text='第一行\n第二行\n\n新段'), '第一行\\\n第二行\n\n新段')

    def test_math_and_chemistry_sources_are_not_escaped(self):
        result = self.md(runs=[{'type': 'text', 'text': '已知 '}, {'type': 'math', 'latex': 'x_1'},
                               {'type': 'text', 'text': '，\n1. 条件'},
                               {'type': 'chemistry', 'latex': r'\mathrm{H_2O}'}])
        self.assertEqual(result, '已知 $x_1$，\\\n1\\. 条件`\\mathrm{H_2O}`')

    def test_html_output_is_unchanged(self):
        self.assertEqual(plain_source({'text': '# a*b'}), '# a*b')

    def test_exported_markdown_file_uses_escaping(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            (root / 'in.json').write_text(json.dumps({'blocks': [
                {'type': 'text', 'text': '# 注意'}, {'type': 'text', 'text': '1. 题目'}]},
                ensure_ascii=False), encoding='utf-8')
            export(root / 'in.json', root / 'out')
            markdown = (root / 'out/result.md').read_text(encoding='utf-8')
            self.assertEqual(markdown, '\\# 注意\n\n1\\. 题目\n')
            self.assertIn('<p># 注意</p>', (root / 'out/result.html').read_text(encoding='utf-8'))


class WordTextCheckTests(unittest.TestCase):
    def test_currency_is_ordinary_text(self):
        for text in ('单价$5，总价$20', '费用 $3 和 $4', 'costs $5 and $20', '$5-$10', 'C:\\Users'):
            check_word_text(text)

    def test_untyped_latex_is_still_refused(self):
        for text in ('$x+1$', '$ x^2 $', '面积为$S$', r'$\frac{1}{2}$', r'\(a\)', r'\sqrt{2}'):
            with self.assertRaisesRegex(ValueError, 'LaTeX found', msg=text):
                check_word_text(text)

    def test_currency_text_exports_to_word(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            (root / 'in.json').write_text(json.dumps({'blocks': [{'type': 'text', 'text': '单价$5，总价$20'}]},
                                                     ensure_ascii=False), encoding='utf-8')
            docx = Document(export_docx(root / 'in.json', root / 'out.docx', word_confirmed=True))
            self.assertEqual(docx.paragraphs[0].text, '单价$5，总价$20')


class FigureSizingTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)

    def test_smallest_valid_canvas_renders_an_accepted_png(self):
        for width, height in ((50, 50), (60, 120), (66, 40 + 10)):
            scene = {'width': width, 'height': height,
                     'elements': [{'id': 'l', 'kind': 'line', 'points': [[5, 5], [width - 5, height - 5]]}]}
            png = self.root / f'{width}x{height}.png'
            render(scene, png.with_suffix('.svg'), png)
            self.assertTrue(check_black_white(png))
            with Image.open(png) as image:
                self.assertGreaterEqual(image.width, 200)
                self.assertGreaterEqual(image.height, 150)
                # Physical size stays 100 scene units per inch after upscaling.
                self.assertAlmostEqual(image.width / image.info['dpi'][0], width / 100, places=2)

    def test_default_scale_is_unchanged(self):
        scene = {'width': 400, 'height': 300,
                 'elements': [{'id': 'l', 'kind': 'line', 'points': [[5, 5], [395, 295]]}]}
        png = self.root / 'default.png'
        render(scene, png.with_suffix('.svg'), png)
        with Image.open(png) as image:
            self.assertEqual(image.size, (1200, 900))
            self.assertAlmostEqual(image.info['dpi'][0], 300, places=0)

    def test_redrawn_figure_uses_png_resolution_in_word(self):
        doc = document(self.root)
        ledger = generate(doc, self.root, self.root / 'images', mode='redraw')
        result, _ = complete(ledger, reviews(ledger), self.root)
        path = self.root / 'reviewed.json'
        path.write_text(json.dumps(result), encoding='utf-8')
        shape = Document(export_docx(path, self.root / 'out.docx', word_confirmed=True)).inline_shapes[0]
        # 400 scene units at 100 units/inch = 4 inches, not the full text width.
        self.assertAlmostEqual(shape.width / Cm(1), 4 * 2.54, delta=0.05)
        self.assertAlmostEqual(shape.width / shape.height, 4 / 3, places=4)

    def test_png_without_dpi_keeps_96_dpi_sizing(self):
        Image.new('RGB', (480, 240), 'white').save(self.root / 'plain.png')
        (self.root / 'in.json').write_text(json.dumps({'blocks': [{'type': 'image', 'path': 'plain.png'}]}),
                                           encoding='utf-8')
        shape = Document(export_docx(self.root / 'in.json', self.root / 'out.docx',
                                     word_confirmed=True)).inline_shapes[0]
        self.assertAlmostEqual(shape.width / Cm(1), 480 / 96 * 2.54, delta=0.05)


class FailurePlaceholderTests(unittest.TestCase):
    def test_placeholder_uses_reader_facing_figure_order(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            doc = document(root, 2)
            doc['figure_task']['figures'][1]['critical_uncertainty'] = True
            ledger = generate(doc, root, root / 'images', mode='redraw')
            result, _ = complete(ledger, reviews(ledger), root)
            texts = [b.get('text') for b in result['blocks'] if b['type'] == 'text' and 'text' in b]
            self.assertIn('[第2张插图：无法可靠重绘]', texts)
            self.assertFalse(any('f2' in (t or '') for t in texts))
            self.assertEqual(result['figure_processing']['failed_figures'], ['f2'])


if __name__ == '__main__':
    unittest.main()
