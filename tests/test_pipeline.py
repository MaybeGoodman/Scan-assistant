"""Pipeline orchestration and prefilled review templates (0.6.0)."""
from copy import deepcopy
import json
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest
from docx import Document
from PIL import Image

TESTS = Path(__file__).resolve().parent
SCRIPTS = TESTS.parent / 'plugins/image-print-extractor/skills/image-print-extractor/scripts'
sys.path.insert(0, str(SCRIPTS))
sys.path.insert(0, str(TESTS))
from test_figures import document  # noqa: E402
from figure_drawing import CHECKS  # noqa: E402
from figure_workflow import complete, generate, review_candidate, review_template  # noqa: E402
from select_content import select_file  # noqa: E402
from export_page import export  # noqa: E402
import pipeline  # noqa: E402

AD = {'category': 'ADVERTISEMENT', 'relation': 'unrelated', 'certainty': 'high', 'reason': '独立推广横幅'}


def write(path, value):
    Path(path).write_text(json.dumps(value, ensure_ascii=False), encoding='utf-8')
    return Path(path)


def read(path):
    return json.loads(Path(path).read_text(encoding='utf-8'))


def fill(state, status='pass', name=None, **check_changes):
    """Simulate the host filling the template after looking at the files."""
    base = Path(state['base'])
    template = read(base / state['files']['review_template'])
    for row in template['reviews']:
        row['status'] = status
        row['checks'] = {check: True for check in CHECKS}
        row['checks'].update(check_changes)
        row['notes'] = ['已对照原图与重绘图']
    return write(base / (name or state['files']['reviews']), template)


class PipelineTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.scan = self.root / 'scan.png'  # the original attachment passed as --source (0.7.0)
        Image.new('RGB', (40, 30), 'white').save(self.scan)

    def text_candidates(self):
        return write(self.root / 'candidates.json', {'blocks': [
            {'type': 'text', 'runs': [{'type': 'text', 'text': '1．已知 '}, {'type': 'math', 'latex': 'a>0'}]},
            {'type': 'text', 'text': '扫码领取资料', 'selection': AD},
            {'type': 'text', 'text': '页脚不确定', 'selection': {'category': 'UNKNOWN', 'relation': 'unknown',
                                                              'certainty': 'uncertain', 'reason': '用途不明'}}]})

    def figure_candidates(self, **kwargs):
        return write(self.root / 'candidates.json', document(self.root, **kwargs))

    def test_text_only_html_matches_manual_steps(self):
        source = self.text_candidates()
        state = pipeline.run_command(source, 'r1', source=self.scan, fmt='html', output=self.root / 'out')
        self.assertEqual(state['status'], 'completed')
        select_file(source, self.root / 'manual.json', self.root / 'manual-report.json')
        export(self.root / 'manual.json', self.root / 'manual-out')
        for name in ('result.html', 'result.md'):
            self.assertEqual((self.root / 'out' / name).read_bytes(), (self.root / 'manual-out' / name).read_bytes())
        self.assertNotIn('扫码领取资料', (self.root / 'out/result.html').read_text(encoding='utf-8'))
        self.assertEqual(state['selection_needs_review'], read(self.root / 'manual-report.json')['needs_review'])
        self.assertIn('页脚不确定', (self.root / 'out/result.html').read_text(encoding='utf-8'))

    def test_docx_requires_word_confirmation(self):
        source = self.text_candidates()
        with self.assertRaisesRegex(ValueError, 'word-confirmed'):
            pipeline.run_command(source, 'r1', source=self.scan, fmt='docx', output=self.root / 'out.docx')
        self.assertFalse((self.root / 'r1-work').exists())
        self.assertFalse((self.root / 'out.docx').exists())
        state = pipeline.run_command(source, 'r2', source=self.scan, fmt='docx', word_confirmed=True, output=self.root / 'out.docx')
        self.assertEqual(state['status'], 'completed')
        self.assertIn('已知', Document(self.root / 'out.docx').paragraphs[0].text)
        self.assertTrue((self.root / 'r2-work/r2.word.review.json').exists())

    def test_figure_task_requires_image_mode(self):
        source = self.figure_candidates()
        with self.assertRaisesRegex(ValueError, 'image-mode'):
            pipeline.run_command(source, 'r1', source=self.scan, fmt='html', output=self.root / 'out')

    def test_ignore_completes_in_one_run(self):
        source = self.figure_candidates()
        state = pipeline.run_command(source, 'r1', source=self.scan, image_mode='ignore', fmt='html', output=self.root / 'out')
        self.assertEqual(state['status'], 'completed')
        html = (self.root / 'out/result.html').read_text(encoding='utf-8')
        self.assertNotIn('<img', html)
        self.assertIn('如图', html)
        self.assertFalse((self.root / 'r1-images').exists())

    def test_preserve_completes_in_one_run(self):
        source = self.figure_candidates()
        state = pipeline.run_command(source, 'r1', source=self.scan, image_mode='preserve', fmt='html', output=self.root / 'out')
        self.assertEqual(state['status'], 'completed')
        self.assertIn('<img', (self.root / 'out/result.html').read_text(encoding='utf-8'))

    def test_redraw_stops_with_null_template_matching_ledger(self):
        source = self.figure_candidates()
        state = pipeline.run_command(source, 'r1', source=self.scan, image_mode='redraw', fmt='html', output=self.root / 'out')
        self.assertEqual(state['status'], 'awaiting_review')
        ledger = read(self.root / state['files']['ledger'])
        template = read(self.root / state['files']['review_template'])
        entry = ledger['records'][0]['attempts'][0]
        row = template['reviews'][0]
        for key in ('source_sha256', 'png_sha256', 'svg_sha256'):
            self.assertEqual(row[key], entry[key])
        self.assertIsNone(row['status'])
        self.assertEqual(set(row['checks']), set(CHECKS))
        self.assertTrue(all(value is None for value in row['checks'].values()))
        self.assertEqual(row['view_files'], {'source': 'source.png', 'png': entry['png_path'], 'svg': entry['svg_path']})
        self.assertEqual(state['view_files'], [row['view_files']])
        self.assertFalse((self.root / 'out').exists())
        self.assertTrue((self.root / 'r1.text-reviewed.json').exists())

    def test_unfilled_template_is_rejected_and_run_stays_resumable(self):
        state = pipeline.run_command(self.figure_candidates(), 'r1', source=self.scan, image_mode='redraw', fmt='html', output=self.root / 'out')
        state_path = self.root / 'r1-work/r1.state.json'
        with self.assertRaisesRegex(ValueError, '模板'):
            pipeline.review_command(state_path, self.root / state['files']['review_template'])
        copy = write(self.root / 'r1-work/copy.json', read(self.root / state['files']['review_template']))
        with self.assertRaisesRegex(ValueError, 'status'):
            pipeline.review_command(state_path, copy)
        saved = read(state_path)
        self.assertEqual(saved['status'], 'awaiting_review')
        self.assertIn('status', saved['error']['message'])
        state = pipeline.review_command(state_path, fill(saved))
        self.assertEqual(state['status'], 'completed')
        self.assertNotIn('error', read(state_path))

    def test_redraw_pass_matches_manual_steps(self):
        source = self.figure_candidates()
        state = pipeline.run_command(source, 'r1', source=self.scan, image_mode='redraw', fmt='html', output=self.root / 'out')
        state = pipeline.review_command(self.root / 'r1-work/r1.state.json', fill(state))
        self.assertEqual(state['status'], 'completed')
        self.assertEqual(state['failed_figures'], [])
        html = (self.root / 'out/result.html').read_text(encoding='utf-8')
        self.assertIn('<img', html)
        self.assertTrue((self.root / 'out/figures/figure-02.svg').exists())
        # Manual route through the unchanged scripts gives the same figure blocks.
        manual_dir = self.root / 'manual'
        manual_dir.mkdir()
        doc = document(manual_dir)
        ledger = generate(doc, manual_dir, manual_dir / 'images', mode='redraw')
        rows = review_template(ledger)
        for row in rows['reviews']:
            row.update(status='pass', checks={c: True for c in CHECKS})
        result, _ = complete(ledger, rows, manual_dir)
        mine = read(self.root / state['files']['figures_reviewed'])
        strip = lambda blocks: [{k: v for k, v in b.items() if k not in {'path', 'svg_path', 'redraw'}} for b in blocks]
        self.assertEqual(strip(mine['blocks']), strip(result['blocks']))

    def test_major_issue_then_retry_passes(self):
        state = pipeline.run_command(self.figure_candidates(), 'r1', source=self.scan, image_mode='redraw', fmt='docx',
                                     word_confirmed=True, output=self.root / 'out.docx')
        state_path = self.root / 'r1-work/r1.state.json'
        state = pipeline.review_command(state_path, fill(state, 'major_issue', labels=False))
        self.assertEqual((state['status'], state['needs_retry']), ('needs_retry', ['f1']))
        fixed = read(self.root / state['files']['figure_source'])
        fixed['figure_task']['figures'][0]['drawing']['scene']['elements'][0]['width'] = 90
        state = pipeline.retry_command(state_path, write(self.root / 'r1.figure-source-fix1.json', fixed))
        self.assertEqual((state['status'], state['attempt']), ('awaiting_review', 2))
        self.assertTrue((self.root / 'r1-work/r1.reviews2.template.json').exists())
        state = pipeline.review_command(state_path, fill(state))
        self.assertEqual(state['status'], 'completed')
        self.assertEqual(len(Document(self.root / 'out.docx').inline_shapes), 1)

    def test_two_major_issues_leave_placeholder(self):
        state = pipeline.run_command(self.figure_candidates(), 'r1', source=self.scan, image_mode='redraw', fmt='html', output=self.root / 'out')
        state_path = self.root / 'r1-work/r1.state.json'
        state = pipeline.review_command(state_path, fill(state, 'major_issue', labels=False))
        fixed = read(self.root / state['files']['figure_source'])
        state = pipeline.retry_command(state_path, write(self.root / 'fix.json', fixed))
        state = pipeline.review_command(state_path, fill(state, 'major_issue', labels=False))
        self.assertEqual(state['status'], 'completed')
        self.assertEqual(state['failed_figures'], ['f1'])
        html = (self.root / 'out/result.html').read_text(encoding='utf-8')
        self.assertIn('[第1张插图：无法可靠重绘]', html)
        self.assertNotIn('<img', html)

    def test_retry_cannot_change_understanding_or_text(self):
        state = pipeline.run_command(self.figure_candidates(), 'r1', source=self.scan, image_mode='redraw', fmt='html', output=self.root / 'out')
        state_path = self.root / 'r1-work/r1.state.json'
        state = pipeline.review_command(state_path, fill(state, 'major_issue', labels=False))
        original = read(self.root / state['files']['figure_source'])
        changed_understanding = deepcopy(original)
        changed_understanding['figure_task']['figures'][0]['understanding']['objects'].append({'id': 'x', 'kind': 'body'})
        changed_text = deepcopy(original)
        changed_text['blocks'][0]['runs'][0]['text'] = '改写的题干'
        for name, value in (('u.json', changed_understanding), ('t.json', changed_text)):
            with self.subTest(name=name), self.assertRaises(ValueError):
                pipeline.retry_command(state_path, write(self.root / name, value))
            self.assertEqual(read(state_path)['status'], 'needs_retry')

    def test_modified_candidates_are_detected(self):
        source = self.figure_candidates()
        state = pipeline.run_command(source, 'r1', source=self.scan, image_mode='redraw', fmt='html', output=self.root / 'out')
        reviews = fill(state)
        data = read(source)
        data['blocks'][0]['runs'][0]['text'] = '被修改'
        write(source, data)
        with self.assertRaisesRegex(ValueError, '被修改'):
            pipeline.review_command(self.root / 'r1-work/r1.state.json', reviews)

    def test_run_id_reuse_and_format(self):
        source = self.text_candidates()
        pipeline.run_command(source, 'r1', source=self.scan, fmt='html', output=self.root / 'out')
        with self.assertRaises(FileExistsError):
            pipeline.run_command(source, 'r1', source=self.scan, fmt='html', output=self.root / 'out2')
        with self.assertRaises(ValueError):
            pipeline.run_command(source, '../x', source=self.scan, fmt='html', output=self.root / 'out3')

    def qr_candidates(self):
        Image.new('RGB', (60, 60), 'black').save(self.root / 'qr.png')
        return write(self.root / 'candidates.json', {
            'blocks': [{'type': 'text', 'id': 'q1', 'text': '1．正文'},
                       {'type': 'image', 'figure_id': 'qr1', 'path': 'qr.png',
                        'selection': {'category': 'QR_PROMOTION', 'relation': 'unrelated', 'certainty': 'high',
                                      'reason': '页脚宣传二维码'}}],
            'figure_task': {'schema_version': 1, 'task_id': 'current-file', 'page_count': 1, 'scan_complete': True,
                            'figures': [{'figure_id': 'qr1', 'source_page': 1, 'source_bbox': [0, 0, 60, 60],
                                         'source_path': 'qr.png', 'parent_block_id': 'q1', 'figure_type': 'qr',
                                         'relation': 'unrelated'}]}})

    def test_explicit_qr_preservation_reaches_export(self):
        source = self.qr_candidates()
        state = pipeline.run_command(source, 'keep', source=self.scan, preserve=['QR_PROMOTION'], image_mode='preserve', fmt='html',
                                     output=self.root / 'kept')
        self.assertEqual(state['status'], 'completed')
        self.assertIn('<img', (self.root / 'kept/result.html').read_text(encoding='utf-8'))
        state = pipeline.run_command(source, 'drop', source=self.scan, image_mode='preserve', fmt='html', output=self.root / 'dropped')
        self.assertEqual(state['figures_removed_by_selection'], ['qr1'])
        self.assertNotIn('<img', (self.root / 'dropped/result.html').read_text(encoding='utf-8'))

    def test_review_candidate_ignores_view_files(self):
        doc = document(self.root)
        ledger = generate(doc, self.root, self.root / 'images', mode='redraw')
        row = review_template(ledger)['reviews'][0]
        row.update(status='pass', checks={c: True for c in CHECKS})
        self.assertIn('view_files', row)
        self.assertEqual(review_candidate(ledger['records'][0]['attempts'][-1], row, self.root), 'pass')

    def test_generate_cli_review_template_is_optional(self):
        source = self.figure_candidates()
        script = str(SCRIPTS / 'figure_workflow.py')
        base = [sys.executable, '-X', 'utf8', script, 'generate', str(source), '--image-mode', 'redraw']
        plain = subprocess.run(base + ['--output-dir', str(self.root / 'a'), '--ledger', str(self.root / 'a.json')],
                               capture_output=True, text=True, encoding='utf-8')
        self.assertEqual(plain.returncode, 0, plain.stderr)
        self.assertEqual(sorted(p.name for p in self.root.glob('*.template.json')), [])
        done = subprocess.run(base + ['--output-dir', str(self.root / 'b'), '--ledger', str(self.root / 'b.json'),
                                      '--review-template', str(self.root / 'b.template.json')],
                              capture_output=True, text=True, encoding='utf-8')
        self.assertEqual(done.returncode, 0, done.stderr)
        self.assertIsNone(read(self.root / 'b.template.json')['reviews'][0]['status'])

    def test_cli_prints_next_action(self):
        source = self.figure_candidates()
        done = subprocess.run([sys.executable, '-X', 'utf8', str(SCRIPTS / 'pipeline.py'), 'run', str(source), '--source', str(self.scan), '--run-id', 'c1',
                               '--image-mode', 'redraw', '--format', 'html', '--output', str(self.root / 'out')],
                              capture_output=True, text=True, encoding='utf-8')
        self.assertEqual(done.returncode, 0, done.stderr)
        result = json.loads(done.stdout)
        self.assertEqual(result['status'], 'awaiting_review')
        self.assertIn('pipeline.py review', result['next_action'])
        status = subprocess.run([sys.executable, '-X', 'utf8', str(SCRIPTS / 'pipeline.py'), 'status', '--state', result['state']],
                                capture_output=True, text=True, encoding='utf-8')
        self.assertEqual(json.loads(status.stdout)['status'], 'awaiting_review')
        bad = subprocess.run([sys.executable, '-X', 'utf8', str(SCRIPTS / 'pipeline.py'), 'run', str(source), '--source', str(self.scan), '--run-id', 'c2',
                              '--format', 'docx', '--output', str(self.root / 'x.docx')],
                             capture_output=True, text=True, encoding='utf-8')
        self.assertEqual(bad.returncode, 1)
        self.assertEqual(json.loads(bad.stdout)['status'], 'error')


if __name__ == '__main__':
    unittest.main()
