import json
import sys
import tempfile
import unittest
from pathlib import Path

SCRIPTS = Path(__file__).resolve().parents[1] / 'plugins/image-print-extractor/skills/image-print-extractor/scripts'
sys.path.insert(0, str(SCRIPTS))
from latex_validation import FormulaError, validate_latex
from export_page import export


class ChemistrySourceTests(unittest.TestCase):
    def test_rejects_typographic_and_plain_text_substitutes(self):
        for source in ('N₂', 'N2', 'SO₄²⁻', r'\mathrm{H2SO4}',
                       r'\mathrm{Ca(OH)2}', r'\mathrm{N₂}', 'Na',
                       r'\mathrm{AgCl}↓'):
            with self.subTest(source=source), self.assertRaises(FormulaError):
                validate_latex(source, chemistry=True)

    def test_preserves_semantics_and_source_spelling(self):
        for source in (r'2\mathrm{N_2}', r'\mathrm{SO_4^{2-}}',
                       r'{}^{23}_{11}\mathrm{Na}', r'\overset{+2}{\mathrm{Fe}}',
                       r'\mathrm{H^+}+\mathrm{OH^-}=\mathrm{H_2O}',
                       r'\mathrm{H_2}+\mathrm{O_2}\rightarrow\mathrm{H_2O}',
                       r'2\mathrm{KClO_3}\xrightarrow[\Delta]{\mathrm{MnO_2}}2\mathrm{KCl}+3\mathrm{O_2}',
                       r'\mathrm{CuSO_4}\cdot 5\mathrm{H_2O}',
                       r'\mathrm{H_2}\xrightarrow{\text{pH7}}\mathrm{H_2}',
                       r'\mathrm{N}_{2}', r'\text{XXX}'):
            self.assertEqual(validate_latex(source, chemistry=True), source)

    def test_html_checks_block_run_and_table_before_writing(self):
        bad = {'type': 'chemistry', 'latex': 'N₂'}
        variants = [bad, {'type': 'text', 'runs': [bad]},
                    {'type': 'table', 'rows': 1, 'cols': 1,
                     'cells': [{'row': 0, 'col': 0, 'runs': [bad]}]}]
        for block in variants:
            with self.subTest(block=block), tempfile.TemporaryDirectory() as temp:
                root = Path(temp)
                source = root / 'input.json'
                source.write_text(json.dumps({'blocks': [block]}), encoding='utf-8')
                with self.assertRaises(FormulaError):
                    export(source, root / 'output')
                self.assertFalse((root / 'output').exists())

    def test_word_invalid_chemistry_is_recorded_including_table(self):
        from docx import Document
        from export_docx import export_docx
        bad = {'type': 'chemistry', 'latex': 'N₂'}
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            source = root / 'input.json'
            source.write_text(json.dumps({'blocks': [
                bad, {'type': 'text', 'runs': [bad]},
                {'type': 'table', 'rows': 1, 'cols': 1,
                 'cells': [{'row': 0, 'col': 0, 'runs': [bad]}]}]}), encoding='utf-8')
            target = export_docx(source, root / 'result.docx', word_confirmed=True)
            document = Document(target)
            self.assertEqual([p.text for p in document.paragraphs], ['XXX', 'XXX'])
            self.assertEqual(document.tables[0].cell(0, 0).text, 'XXX')
            report = json.loads(target.with_suffix('.review.json').read_text(encoding='utf-8'))
            self.assertEqual(report['needs_review'], 3)
            self.assertEqual(report['expressions'][2]['location'], 'blocks[2].cells[0].runs[0]')
            self.assertTrue(all(x['latex'] == 'N₂' for x in report['expressions']))

    def test_ordinary_prose_is_not_classified_as_chemistry(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            source = root / 'input.json'
            source.write_text(json.dumps({'blocks': [{'type': 'text', 'text': '型号 N2，2026 年'}]}), encoding='utf-8')
            output = export(source, root / 'output')
            self.assertIn('型号 N2，2026 年', output.read_text(encoding='utf-8'))


if __name__ == '__main__':
    unittest.main()
