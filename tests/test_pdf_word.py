import json
import sys
import tempfile
import unittest
from pathlib import Path
from zipfile import ZipFile
from docx import Document
from PIL import Image
from reportlab.pdfgen import canvas

scripts = Path(__file__).resolve().parents[1] / 'plugins/image-print-extractor/skills/image-print-extractor/scripts'
sys.path.insert(0, str(scripts))
from export_docx import export_docx, paragraph_layout
from prepare_pdf import prepare


class PdfWordTests(unittest.TestCase):
    def test_question_figure_and_options_keep_together_in_actual_docx(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            Image.new('RGB', (300, 100), 'white').save(root / 'figure.png')
            blocks = [{'type':'text','text':'11．如图所示，选择正确选项。'},
                      {'type':'image','path':'figure.png'},
                      *[{'type':'text','text':letter+'．选项'} for letter in 'ABCD'],
                      {'type':'text','text':'12．另一道题。'},
                      {'type':'image','path':'figure.png'},
                      {'type':'text','text':'独立正文。'}]
            source = root / 'reviewed.json'
            source.write_text(json.dumps({'blocks':blocks}), encoding='utf-8')
            doc = Document(export_docx(source, root/'out.docx', word_confirmed=True))
            self.assertEqual([p.paragraph_format.keep_with_next for p in doc.paragraphs],
                             [True, True, True, True, True, False, True, False, False])
            self.assertTrue(all(p.paragraph_format.keep_together for p in doc.paragraphs))
            self.assertEqual([p.text for p in doc.paragraphs if p.text],
                             [b['text'] for b in blocks if b['type']=='text'])

    def test_pagination_does_not_link_unrelated_blocks(self):
        blocks = [{'type':'text','text':'普通正文'}, {'type':'text','text':'后续正文'},
                  {'type':'text','text':'A．独立编号'}, {'type':'text','text':'C．不连续编号'},
                  {'type':'text','text':'1．列表条目'}, {'type':'text','text':'2．另一个条目'}]
        self.assertFalse(any(paragraph_layout(blocks, i)['keep_with_next'] for i in range(len(blocks))))

    def test_explicit_pagination_override_and_validation(self):
        blocks = [{'type':'text','text':'1．题干','layout':{'keep_with_next':False,'page_break_before':True}},
                  {'type':'image','path':'figure.png'}]
        self.assertFalse(paragraph_layout(blocks, 0)['keep_with_next'])
        self.assertTrue(paragraph_layout(blocks, 0)['page_break_before'])
        for invalid in ({'keep_with_next':'false'}, {'unknown':True}, []):
            with self.subTest(invalid=invalid), self.assertRaises(ValueError):
                paragraph_layout([{'type':'text','text':'正文','layout':invalid}], 0)

    def test_confirmation_required(self):
        with self.assertRaisesRegex(ValueError, 'Confirm Word'):
            export_docx('unused.json', 'unused.docx')

    def test_editable_content_merges_and_embedded_png(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            Image.new('RGB', (300, 100), 'white').save(root / 'figure.png')
            source = root / 'reviewed.json'
            blocks = [
                {'type':'text', 'runs':[{'type':'text','text':'跨页续句已接回，公式 '},
                    {'type':'math','latex':'x^2+1'}, {'type':'text','text':' 保留。'}]},
                {'type':'table', 'rows':3, 'cols':3, 'cells':[
                    {'row':0, 'col':0, 'rowspan':2, 'colspan':2, 'text':'合并表头'},
                    {'row':0, 'col':2, 'text':'数值'},
                    {'row':1, 'col':2, 'text':''},
                    {'row':2, 'col':0, 'text':'A'},
                    {'row':2, 'col':1, 'text':'XXX'},
                    {'row':2, 'col':2, 'runs':[{'type':'math','latex':'a>0'}]}]},
                {'type':'image', 'path':'figure.png', 'alt':'印刷插图'},
                {'type':'text', 'text':'与正文相关的来源说明保留。'}]
            source.write_text(json.dumps({'blocks':blocks}), encoding='utf-8')
            target = export_docx(source, root / 'result.docx', word_confirmed=True)
            doc = Document(target)
            self.assertEqual(doc.paragraphs[0].text, '跨页续句已接回，公式  保留。')
            self.assertEqual(doc.paragraphs[-1].text, blocks[-1]['text'])
            table = doc.tables[0]
            self.assertEqual(table.cell(0, 0).text, '合并表头')
            self.assertEqual(table.cell(1, 1).text, '合并表头')
            self.assertEqual(table.cell(1, 2).text, '')
            self.assertEqual(len(doc.inline_shapes), 1)
            tags = [el.tag.rsplit('}',1)[-1] for el in doc.element.body]
            self.assertEqual(tags, ['p','tbl','p','p','sectPr'])
            with ZipFile(target) as z:
                self.assertTrue(any(name.startswith('word/media/') for name in z.namelist()))
                xml = z.read('word/document.xml').decode('utf-8')
                self.assertIn('w:gridSpan', xml)
                self.assertIn('w:vMerge', xml)
                self.assertIn('m:oMath', xml)
                self.assertNotIn('$x^2+1$', xml)
            with self.assertRaises(FileExistsError):
                export_docx(source, target, word_confirmed=True)

    def test_invalid_table_leaves_no_file(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            source = root / 'bad.json'
            source.write_text(json.dumps({'blocks':[{'type':'table','rows':1,'cols':1,'cells':[]}]}))
            with self.assertRaises(ValueError):
                export_docx(source, root / 'bad.docx', word_confirmed=True)
            self.assertFalse((root / 'bad.docx').exists())

    def test_pdf_all_pages_and_unfiltered_candidates(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            pdf = canvas.Canvas(str(root / 'source.pdf'))
            pdf.drawString(70,700,'PRINTED CONTENT')
            pdf.drawString(70,600,'DOWNLOAD WATERMARK')
            pdf.showPage()
            pdf.showPage()  # Blank/textless page still needs visual review.
            pdf.save()
            result = prepare(root / 'source.pdf', root / 'pages', 72)
            self.assertEqual(result['page_count'], 2)
            self.assertEqual([p['page'] for p in result['pages']], [1, 2])
            self.assertIn('PRINTED CONTENT', result['pages'][0]['text_candidate'])
            self.assertIn('DOWNLOAD WATERMARK', result['pages'][0]['text_candidate'])
            self.assertEqual(result['pages'][1]['text_candidate'], '')
            for page in result['pages']:
                with Image.open(root / 'pages' / page['image']) as image:
                    self.assertEqual(image.format, 'PNG')
            with self.assertRaises(FileExistsError):
                prepare(root / 'source.pdf', root / 'pages')

if __name__ == '__main__':
    unittest.main()
